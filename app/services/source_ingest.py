from __future__ import annotations

import ipaddress
import mimetypes
import socket
from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib.parse import urljoin, urlsplit

import httpx
from docx import Document
from pypdf import PdfReader

MAX_SOURCE_BYTES = 14 * 1024 * 1024
MAX_REDIRECTS = 3
_AUDIO_MIME_TYPES = {
    ".aac": "audio/aac",
    ".flac": "audio/flac",
    ".m4a": "audio/mp4",
    ".mp3": "audio/mpeg",
    ".ogg": "audio/ogg",
    ".opus": "audio/opus",
    ".wav": "audio/wav",
    ".webm": "audio/webm",
}


@dataclass
class SourceContent:
    text: str | None = None
    media: bytes | None = None
    mime_type: str | None = None


def is_audio_filename(filename: str) -> bool:
    return Path(filename).suffix.lower() in _AUDIO_MIME_TYPES


def mime_type_for_audio(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in _AUDIO_MIME_TYPES:
        raise ValueError("Choose an MP3, WAV, M4A, AAC, OGG, FLAC, OPUS, or WebM audio file.")
    return _AUDIO_MIME_TYPES[suffix]


def extract_file_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        text = data.decode("utf-8-sig")
    elif suffix == ".docx":
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        paragraphs.extend(
            "\t".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
        text = "\n".join(paragraphs)
    elif suffix == ".pdf":
        reader = PdfReader(BytesIO(data))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    else:
        raise ValueError("Supported documents are PDF, DOCX, TXT, and Markdown; audio files are also accepted.")

    text = text.strip()
    if not text and suffix != ".pdf":
        raise ValueError("No readable text was found in that file.")
    return text


def fetch_url_source(url: str, user_agent: str = "PidginLaw/1.0") -> SourceContent:
    current_url = url
    with httpx.Client(timeout=20.0, follow_redirects=False, headers={"User-Agent": user_agent}) as client:
        for redirect_count in range(MAX_REDIRECTS + 1):
            _validate_public_url(current_url)
            with client.stream("GET", current_url) as response:
                if response.is_redirect:
                    if redirect_count == MAX_REDIRECTS:
                        raise ValueError("That link redirects too many times.")
                    location = response.headers.get("location")
                    if not location:
                        raise ValueError("That link returned an invalid redirect.")
                    current_url = urljoin(current_url, location)
                    continue

                response.raise_for_status()
                content_length = response.headers.get("content-length")
                if content_length and int(content_length) > MAX_SOURCE_BYTES:
                    raise ValueError("That link is too large. The maximum supported size is 14 MB.")
                body = bytearray()
                for chunk in response.iter_bytes():
                    body.extend(chunk)
                    if len(body) > MAX_SOURCE_BYTES:
                        raise ValueError("That link is too large. The maximum supported size is 14 MB.")
                mime_type = response.headers.get("content-type", "").split(";", 1)[0].strip().lower()
                guessed_type, _encoding = mimetypes.guess_type(urlsplit(current_url).path)
                mime_type = mime_type or guessed_type or "application/octet-stream"
                data = bytes(body)

                if mime_type == "application/pdf" or (
                    mime_type == "application/octet-stream"
                    and urlsplit(current_url).path.lower().endswith(".pdf")
                ):
                    return SourceContent(media=data, mime_type="application/pdf")
                if mime_type.startswith("audio/") or _audio_mime_from_url(current_url):
                    return SourceContent(media=data, mime_type=mime_type if mime_type.startswith("audio/") else _audio_mime_from_url(current_url))
                if mime_type in {"text/plain", "text/markdown"}:
                    return SourceContent(text=data.decode(response.encoding or "utf-8", errors="replace").strip())
                if mime_type in {"text/html", "application/xhtml+xml"}:
                    parser = _ArticleTextParser()
                    parser.feed(data.decode(response.encoding or "utf-8", errors="replace"))
                    text = parser.text()
                    if not text:
                        raise ValueError("No readable article text was found at that link.")
                    return SourceContent(text=text)
                raise ValueError("That link must point to a public article, PDF, or audio file.")

    raise ValueError("Could not fetch that link.")


def _audio_mime_from_url(url: str) -> str | None:
    return _AUDIO_MIME_TYPES.get(Path(urlsplit(url).path).suffix.lower())


def _validate_public_url(url: str) -> None:
    parsed = urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Enter a public HTTP or HTTPS link.")

    host = parsed.hostname.rstrip(".").lower()
    if host == "localhost" or host.endswith((".localhost", ".local")):
        raise ValueError("Private and local-network links are not allowed.")

    try:
        addresses = [ipaddress.ip_address(host)]
    except ValueError:
        try:
            records = socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80))
        except socket.gaierror as exc:
            raise ValueError("That link could not be resolved.") from exc
        addresses = [ipaddress.ip_address(record[4][0].split("%", 1)[0]) for record in records]

    if not addresses or any(not address.is_global for address in addresses):
        raise ValueError("Private and local-network links are not allowed.")


class _ArticleTextParser(HTMLParser):
    _SKIP_TAGS = {"script", "style", "nav", "header", "footer", "aside", "noscript", "svg"}
    _BLOCK_TAGS = {"article", "br", "div", "h1", "h2", "h3", "li", "main", "p", "section"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip_depth = 0
        self._chunks: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1
        elif not self._skip_depth and tag in self._BLOCK_TAGS:
            self._chunks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP_TAGS and self._skip_depth:
            self._skip_depth -= 1
        elif not self._skip_depth and tag in self._BLOCK_TAGS:
            self._chunks.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self._chunks.append(data)

    def text(self) -> str:
        lines = [" ".join(line.split()) for line in "".join(self._chunks).splitlines()]
        return "\n\n".join(line for line in lines if line)